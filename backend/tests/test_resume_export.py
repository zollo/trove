"""Converting resumes to PDF / DOCX for download."""
from __future__ import annotations

import io

import docx

from app.core import resume_export, resume_text
from app.core.resume_export import ExportFormat, Run

MARKDOWN = """# Jane Doe
jane@example.com
Denver, CO

## Experience

### Staff Engineer, Acme
- Cut deploy time by **80%** with *zero* downtime
- Built [Trove](https://example.com) using `FastAPI`

---

Plain closing paragraph.
"""


def test_parse_markdown():
    blocks = resume_export.parse_markdown(MARKDOWN)
    assert [b.kind for b in blocks] == ["h1", "p", "h2", "h3", "li", "li", "hr", "p"]
    # Consecutive lines keep their line breaks.
    assert blocks[1].runs == [Run("jane@example.com\nDenver, CO")]
    assert blocks[4].runs == [
        Run("Cut deploy time by "),
        Run("80%", bold=True),
        Run(" with "),
        Run("zero", italic=True),
        Run(" downtime"),
    ]
    assert "".join(r.text for r in blocks[5].runs) == "Built Trove using FastAPI"


def test_parse_html():
    html = "<h1>Jane</h1><p>Loves <b>Python</b></p><ul><li>One</li></ul><script>x()</script>"
    blocks = resume_export.parse_html(html)
    assert [b.kind for b in blocks] == ["h1", "p", "li"]
    assert Run("Python", bold=True) in blocks[1].runs


def test_heading_keeps_literal_hash():
    blocks = resume_export.parse_markdown("## C#\n## .NET / C# ##\n")
    assert [b.runs[0].text for b in blocks] == ["C#", ".NET / C#"]


def test_html_table_cells_stay_apart():
    html = "<table><tr><td>2020-2024</td><td>Acme</td><td>Staff Engineer</td></tr></table>"
    (row,) = resume_export.parse_html(html)
    assert "".join(r.text for r in row.runs) == "2020-2024 | Acme | Staff Engineer"


def test_docx_keeps_body_order_and_hyperlinks():
    from docx.oxml.parser import OxmlElement

    document = docx.Document()
    contact = document.add_paragraph("Email: ")
    link = OxmlElement("w:hyperlink")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "jane@example.com"
    run.append(text)
    link.append(run)
    contact._p.append(link)
    document.add_heading("Experience", level=2)
    document.add_table(rows=1, cols=2).rows[0].cells[0].text = "Acme"
    document.add_heading("Education", level=2)
    buf = io.BytesIO()
    document.save(buf)

    blocks = resume_export.parse_docx(buf.getvalue())
    texts = ["".join(r.text for r in b.runs) for b in blocks]
    assert texts == ["Email: jane@example.com", "Experience", "Acme | ", "Education"]


def test_markdown_round_trips_through_docx_and_pdf():
    name, data = resume_export.convert("cv.md", MARKDOWN.encode(), None, ExportFormat.docx)
    assert name == "cv.docx"
    document = docx.Document(io.BytesIO(data))
    assert document.paragraphs[0].text == "Jane Doe"
    assert any(r.bold and r.text == "80%" for p in document.paragraphs for r in p.runs)

    name, pdf = resume_export.convert(name, data, None, ExportFormat.pdf)
    assert name == "cv.pdf" and pdf.startswith(b"%PDF")
    text = resume_text.extract_text(name, pdf)
    assert "Staff Engineer" in text and "Cut deploy time" in text


def test_same_format_is_unchanged():
    assert resume_export.convert("cv.PDF", b"%PDF-1.4", None, ExportFormat.pdf) == (
        "cv.PDF",
        b"%PDF-1.4",
    )


def test_pdf_falls_back_to_core_fonts(monkeypatch):
    monkeypatch.setattr(resume_export, "_FONT_DIRS", ())
    text = "Zoë — “quoted” • 東京".encode()
    _, pdf = resume_export.convert("cv.txt", text, None, ExportFormat.pdf)
    assert "Zoë" in resume_text.extract_text("x.pdf", pdf)
