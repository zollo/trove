"""Resume file validation and plain-text extraction."""
from __future__ import annotations

import io
import logging
import mimetypes
import os

logger = logging.getLogger("trove")

# Extension -> canonical content type.
ALLOWED_TYPES: dict[str, str] = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".doc": "application/msword",
    ".odt": "application/vnd.oasis.opendocument.text",
    ".rtf": "application/rtf",
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".html": "text/html",
}

# Characters kept in a sanitised download filename.
_SAFE = set("abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789-_. ()")


def extension(filename: str) -> str:
    return os.path.splitext(filename)[1].lower()


def is_allowed(filename: str) -> bool:
    return extension(filename) in ALLOWED_TYPES


def content_type_for(filename: str) -> str:
    ext = extension(filename)
    return ALLOWED_TYPES.get(ext) or mimetypes.guess_type(filename)[0] or "application/octet-stream"


def safe_filename(filename: str, default: str = "resume") -> str:
    name = os.path.basename(filename.replace("\\", "/")).strip()
    name = "".join(c for c in name if c in _SAFE).strip(". ")
    return name[:200] or default


def extract_text(filename: str, data: bytes) -> str | None:
    """Best-effort text extraction. Returns None for unsupported or broken files."""
    ext = extension(filename)
    try:
        if ext in (".txt", ".md", ".html"):
            return data.decode("utf-8", errors="replace").strip()
        if ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()
        if ext == ".docx":
            import docx

            document = docx.Document(io.BytesIO(data))
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(parts).strip()
    except Exception:  # noqa: BLE001
        logger.warning("Text extraction failed for %s", filename, exc_info=True)
    return None
