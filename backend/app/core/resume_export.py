"""Convert a stored resume to PDF or DOCX for download.

Resumes arrive in many formats (AI headhunters usually upload Markdown). Each
source is parsed into a small block model (headings, paragraphs, bullets with
bold/italic runs) and rendered to the requested format. A file that is already
in the requested format is served unchanged.
"""
from __future__ import annotations

import enum
import io
import logging
import os
import re
from dataclasses import dataclass, field
from html.parser import HTMLParser

from app.core import resume_text

logger = logging.getLogger("trove")


class ExportFormat(enum.StrEnum):
    pdf = "pdf"
    docx = "docx"


class ConversionError(Exception):
    """The source file can't be converted to the requested format."""


@dataclass
class Run:
    text: str
    bold: bool = False
    italic: bool = False


@dataclass
class Block:
    # "h1".."h3", "p", "li" or "hr"
    kind: str
    runs: list[Run] = field(default_factory=list)


# ── Parsing ─────────────────────────────────────────────────────────────

_INLINE = re.compile(
    r"\*\*(?P<b1>.+?)\*\*|__(?P<b2>.+?)__"
    r"|(?<![\w*])\*(?P<i1>[^*\s][^*]*?)\*(?![\w*])|(?<!\w)_(?P<i2>[^_\s][^_]*?)_(?!\w)"
    r"|`(?P<code>[^`]+)`|!?\[(?P<link>[^\]]*)\]\([^)]*\)"
)
# Closing hashes need a space before them, so "## C#" keeps its hash.
_HEADING = re.compile(r"^(#{1,6})\s+(.*?)(?:\s+#+)?\s*$")
_BULLET = re.compile(r"^\s*(?:[-*+•]|\d+[.)])\s+(.*)$")
_RULE = re.compile(r"^\s*([-*_])(\s*\1){2,}\s*$")


def _inline(text: str, bold: bool = False) -> list[Run]:
    """Markdown emphasis, code and links -> runs (one level of nesting is plenty)."""
    runs: list[Run] = []
    pos = 0
    for m in _INLINE.finditer(text):
        if m.start() > pos:
            runs.append(Run(text[pos : m.start()], bold=bold))
        if m["b1"] is not None or m["b2"] is not None:
            runs.extend(
                Run(r.text, bold=True, italic=r.italic) for r in _inline(m["b1"] or m["b2"])
            )
        elif m["i1"] is not None or m["i2"] is not None:
            runs.append(Run(m["i1"] or m["i2"], bold=bold, italic=True))
        else:
            runs.append(Run(m["code"] if m["code"] is not None else m["link"], bold=bold))
        pos = m.end()
    if pos < len(text):
        runs.append(Run(text[pos:], bold=bold))
    return runs


def parse_markdown(text: str) -> list[Block]:
    blocks: list[Block] = []
    para: list[str] = []

    def flush() -> None:
        if para:
            # Keep the author's line breaks: resumes put contact details on
            # consecutive lines and expect them to stay that way.
            blocks.append(Block("p", _inline("\n".join(para))))
            para.clear()

    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            flush()
        elif m := _HEADING.match(line):
            flush()
            level = min(len(m[1]), 3)
            blocks.append(Block(f"h{level}", _inline(m[2])))
        elif _RULE.match(line):
            flush()
            blocks.append(Block("hr"))
        elif m := _BULLET.match(line):
            flush()
            blocks.append(Block("li", _inline(m[1].strip())))
        else:
            para.append(line.strip())
    flush()
    return blocks


def parse_plain(text: str) -> list[Block]:
    """Plain text (or text pulled out of a PDF): one paragraph per line."""
    blocks: list[Block] = []
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        if m := _BULLET.match(line):
            blocks.append(Block("li", [Run(m[1].strip())]))
        else:
            blocks.append(Block("p", [Run(line)]))
    return blocks


class _HTMLBlocks(HTMLParser):
    _BLOCK = {"p", "div", "li", "h1", "h2", "h3", "h4", "h5", "h6", "tr", "section"}
    _SKIP = {"script", "style", "head", "title"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.blocks: list[Block] = []
        self.current: Block | None = None
        self.bold = 0
        self.italic = 0
        self.skip = 0

    def _close(self) -> None:
        if self.current and "".join(r.text for r in self.current.runs).strip():
            first = self.current.runs[0]
            first.text = first.text.lstrip()
            self.blocks.append(self.current)
        self.current = None

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag in self._SKIP:
            self.skip += 1
        elif tag in self._BLOCK:
            self._close()
            kind = "li" if tag == "li" else "p"
            if tag[0] == "h" and tag[1:].isdigit():
                kind = f"h{min(int(tag[1:]), 3)}"
            self.current = Block(kind)
        elif tag == "hr":
            self._close()
            self.blocks.append(Block("hr"))
        elif tag in ("td", "th") and self.current and self.current.runs:
            # Cells of a row share one block; keep them apart.
            self.current.runs.append(Run(" | "))
        elif tag == "br" and self.current:
            self.current.runs.append(Run("\n"))
        elif tag in ("b", "strong"):
            self.bold += 1
        elif tag in ("i", "em"):
            self.italic += 1

    def handle_endtag(self, tag: str) -> None:
        if tag in self._SKIP:
            self.skip = max(0, self.skip - 1)
        elif tag in self._BLOCK:
            self._close()
        elif tag in ("b", "strong"):
            self.bold = max(0, self.bold - 1)
        elif tag in ("i", "em"):
            self.italic = max(0, self.italic - 1)

    def handle_data(self, data: str) -> None:
        if self.skip:
            return
        text = re.sub(r"\s+", " ", data)
        if not text.strip() and not self.current:
            return
        if self.current is None:
            self.current = Block("p")
        self.current.runs.append(Run(text, bold=bool(self.bold), italic=bool(self.italic)))


def parse_html(text: str) -> list[Block]:
    parser = _HTMLBlocks()
    parser.feed(text)
    parser.close()
    parser._close()
    return parser.blocks


def parse_docx(data: bytes) -> list[Block]:
    import docx
    from docx.table import Table
    from docx.text.hyperlink import Hyperlink

    document = docx.Document(io.BytesIO(data))
    blocks: list[Block] = []
    # Body order, so tables stay under the section they belong to.
    for item in document.iter_inner_content():
        if isinstance(item, Table):
            for row in item.rows:
                blocks.append(Block("p", [Run(" | ".join(cell.text for cell in row.cells))]))
            continue
        if not item.text.strip():
            continue
        style = (item.style.name if item.style is not None else "") or ""
        if style == "Title":
            kind = "h1"
        elif style.startswith("Heading"):
            level = style.removeprefix("Heading").strip()
            kind = f"h{min(int(level), 3)}" if level.isdigit() else "h2"
        elif "List" in style:
            kind = "li"
        else:
            kind = "p"
        runs: list[Run] = []
        # Hyperlinks (often the email or LinkedIn line) hold runs of their own.
        for part in item.iter_inner_content():
            for r in part.runs if isinstance(part, Hyperlink) else [part]:
                if r.text:
                    runs.append(Run(r.text, bold=bool(r.bold), italic=bool(r.italic)))
        blocks.append(Block(kind, runs or [Run(item.text)]))
    return blocks


def parse(filename: str, data: bytes, text_content: str | None) -> list[Block]:
    ext = resume_text.extension(filename)
    if ext == ".md":
        return parse_markdown(data.decode("utf-8", errors="replace"))
    if ext == ".html":
        return parse_html(data.decode("utf-8", errors="replace"))
    if ext == ".docx":
        return parse_docx(data)
    if ext == ".txt":
        return parse_plain(data.decode("utf-8", errors="replace"))
    if text_content:
        # PDF (and anything else we extracted text from at upload time).
        return parse_plain(text_content)
    raise ConversionError(f"Can't convert a {ext or 'file without an extension'} resume")


# ── Rendering ───────────────────────────────────────────────────────────


def render_docx(blocks: list[Block]) -> bytes:
    import docx
    from docx.shared import Pt

    document = docx.Document()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for block in blocks:
        if block.kind == "hr":
            document.add_paragraph()
            continue
        if block.kind.startswith("h"):
            para = document.add_heading(level=int(block.kind[1]))
        elif block.kind == "li":
            para = document.add_paragraph(style="List Bullet")
        else:
            para = document.add_paragraph()
        for run in block.runs:
            # python-docx turns "\n" into a line break.
            out = para.add_run(run.text)
            out.bold = run.bold or None
            out.italic = run.italic or None
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


# DejaVu covers far more than the PDF core fonts' Latin-1. The backend image
# installs fonts-dejavu-core; elsewhere we fall back to Helvetica.
_FONT_DIRS = ("/usr/share/fonts/truetype/dejavu", "/usr/share/fonts/dejavu")
# Typographic characters resumes use a lot, mapped for the Latin-1 fallback.
_LATIN1 = str.maketrans({"‘": "'", "’": "'", "“": '"', "”": '"', "–": "-", "—": "-",
                         "•": "-", "…": "...", " ": " ", "→": "->"})


def _find_font() -> tuple[str, str] | None:
    for d in _FONT_DIRS:
        regular, bold = os.path.join(d, "DejaVuSans.ttf"), os.path.join(d, "DejaVuSans-Bold.ttf")
        if os.path.exists(regular) and os.path.exists(bold):
            return regular, bold
    return None


def render_pdf(blocks: list[Block]) -> bytes:
    from fpdf import FPDF

    pdf = FPDF(format="letter")
    pdf.set_margins(18, 16, 18)
    pdf.set_auto_page_break(True, margin=16)
    pdf.add_page()

    font = _find_font()
    if font:
        family = "DejaVu"
        regular, bold = font
        # fonts-dejavu-core ships no oblique Sans, so italic text stays upright.
        for style, path in (("", regular), ("B", bold), ("I", regular), ("BI", bold)):
            pdf.add_font(family, style, path)

        def clean(text: str) -> str:
            return text
    else:
        family = "Helvetica"

        def clean(text: str) -> str:
            return text.translate(_LATIN1).encode("latin-1", "replace").decode("latin-1")

    sizes = {"h1": 18, "h2": 13, "h3": 11.5, "p": 10.5, "li": 10.5}
    left = pdf.l_margin
    for block in blocks:
        if block.kind == "hr":
            pdf.ln(2)
            pdf.set_draw_color(180, 180, 180)
            pdf.line(left, pdf.get_y(), pdf.w - pdf.r_margin, pdf.get_y())
            pdf.ln(3)
            continue
        size = sizes[block.kind]
        heading = block.kind.startswith("h")
        line_h = size * 0.5
        if heading:
            pdf.ln(3 if block.kind != "h1" else 0)
        if block.kind == "li":
            pdf.set_font(family, "", size)
            pdf.set_x(left + 3)
            pdf.write(line_h, clean("•"))
            pdf.set_left_margin(left + 8)
            pdf.set_x(left + 8)
        for run in block.runs:
            style = ("B" if run.bold or heading else "") + ("I" if run.italic else "")
            pdf.set_font(family, style, size)
            pdf.write(line_h, clean(run.text))
        pdf.set_left_margin(left)
        pdf.ln(line_h)
        if block.kind == "h2":
            pdf.set_draw_color(180, 180, 180)
            pdf.line(left, pdf.get_y() + 0.5, pdf.w - pdf.r_margin, pdf.get_y() + 0.5)
            pdf.ln(2)
        else:
            pdf.ln(1.5 if block.kind != "li" else 0.8)
    return bytes(pdf.output())


def convert(
    filename: str, data: bytes, text_content: str | None, fmt: ExportFormat
) -> tuple[str, bytes]:
    """Return (filename, bytes) of the resume in ``fmt``; unchanged if it already is."""
    stem, ext = os.path.splitext(filename)
    target = f".{fmt.value}"
    if ext.lower() == target:
        return filename, data
    try:
        blocks = parse(filename, data, text_content)
    except ConversionError:
        raise
    except Exception as exc:  # noqa: BLE001 - a corrupt upload, most likely
        logger.warning("Could not parse %s for export", filename, exc_info=True)
        raise ConversionError("The resume file could not be read for conversion") from exc
    if not blocks:
        raise ConversionError("The resume has no text to convert")
    out = render_pdf(blocks) if fmt is ExportFormat.pdf else render_docx(blocks)
    return f"{stem or 'resume'}{target}", out
